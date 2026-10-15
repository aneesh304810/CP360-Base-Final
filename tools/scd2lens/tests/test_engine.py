import json

from scd2lens.report.json_ import render_json
from scd2lens.report.markdown import render_markdown
from scd2lens.report.xlsx import write_xlsx
from scd2lens.rules.engine import build_card


def test_dim_account_card_offline(export, settings):
    c = build_card(export, "DIM_ACCOUNT", settings)
    assert c.writers == ["SCN_DIM_ACCOUNT", "SCN_DIM_MASTER_ACCOUNT"]
    assert c.surrogate_key == "ACCOUNT_KEY"
    assert c.natural_key == ["ACCOUNT_NUMBER"] and c.version_key == ["ACCOUNT_NUMBER", "START_DATE"]
    assert c.partition and c.partition.column == "MASTER_IND" and c.partition.values == {"N": "SCN_DIM_ACCOUNT", "Y": "SCN_DIM_MASTER_ACCOUNT"}
    texts = {r.id: r for r in c.rules}
    by_text = " || ".join(r.text for r in c.rules)
    assert "START_DATE = EOD → PRIOR_BUSINESS_DATE, INTRADAY → CURRENT_BUSINESS_DATE from DIM_CONTROL_DATES" in by_text
    assert "END_DATE = TRUNC(" in by_text and "-1 day" in by_text and "when the key appears in the current batch" in by_text
    assert "Current row = END_DATE IS NULL; ACTIVE_IND is not maintained by end-dating" in by_text
    assert "Same-day reload overwrites the version" in by_text
    assert "SCD1 overrides on the open row: DATE_INC_CASH_BECAME_OD, DATE_PRIN_CASH_BECAME_OD, TRANS_TO_PRIN_DATE" in by_text
    assert "A second scenario closes versions on the same table [MASTER_IND = 'Y']" in by_text
    assert all(r.confidence == "PROVEN" for r in c.rules), [(r.id, r.confidence) for r in c.rules]
    assert all(r.evidence for r in c.rules)
    ev = " || ".join(r.evidence for r in c.rules)
    assert "SCN_DIM_ACCOUNT step 14 task 10" in ev and "PRC_ACCOUNT_EDM_END_DATE" in ev and "task 420" in ev
    assert any("task 50" in q for q in c.open_questions), c.open_questions
    assert c.catalog_mapping["ACCOUNT_KEY"] == "ENTITY_SK" and c.catalog_mapping["ACCOUNT_NUMBER"] == "BUSINESS_KEY"
    assert c.catalog_mapping["ACTIVE_IND"] == "IS_CURRENT (unreliable)" and c.catalog_mapping["END_DATE"] == "VALID_TO_TS"


def test_runtime_evidence_adds_inferred_rule(export, settings):
    rt = {"dup_summary": {"total_rows": 272373, "dup_rows": 275, "dup_pct": 0.1, "dup_keys": 3},
          "windows": {"bad_windows": 0, "gaps": 0, "overlaps": 0}, "open_rows": {"keys_many_open": 0, "keys_none_open": 12},
          "active_ind": {"open_not_flagged": 4000, "closed_but_flagged": 0}, "freshness": {"alert": None},
          "volume_alerts": ["2026-03: 5 versions vs trailing median 400"], "dup_by_month": [{"month": "2026-02", "versions": 409, "dup_rows": 20}]}
    c = build_card(export, "DIM_ACCOUNT", settings, runtime=rt)
    r = [x for x in c.rules if "real change" in x.text]
    assert r and r[0].confidence == "INFERRED" and "0.1%" in r[0].evidence and "truncated" in r[0].evidence
    assert c.quality.rows == 272373 and c.quality.dup_keys == 3 and c.quality.volume_alerts


def test_reports_render(export, settings, tmp_path):
    c = build_card(export, "DIM_ACCOUNT", settings)
    md = render_markdown(c)
    assert "# DIM_ACCOUNT — SCD2 rule card" in md and "| R1 |" in md and "PRC_ACCOUNT_EDM_END_DATE" in md and "## Open questions" in md
    j = json.loads(render_json(c))
    assert j["pattern"] == "MERGE_SCD2" and j["partition"]["values"]["Y"] == "SCN_DIM_MASTER_ACCOUNT"
    p = write_xlsx([c, build_card(export, "DIM_IP_RELATIONSHIP", settings)], tmp_path / "r.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["cards", "rules", "quality", "dup_by_month", "open_questions", "catalog_mapping"]
    assert wb["cards"].max_row == 3 and wb["cards"]["A1"].fill.fgColor.rgb.endswith("10193B")
