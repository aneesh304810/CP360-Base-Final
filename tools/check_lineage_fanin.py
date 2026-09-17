"""How much fan-in is the current lineage_id discarding? Read-only.

WHY THE DATABASE CANNOT ANSWER THIS

The old lineage_id is "{dwh_target_table}:{dwh_target_column}" and load()
upserts on it. legacy_lineage's grain is (target column x source), so a DWH
column fed by two source chains is TWO rows in the sheet with ONE id: the
second overwrites the first at load time.

So the loss happens before anything reaches Oracle. Querying legacy_lineage
afterwards shows one row per target column and looks perfectly consistent —
it cannot show you what it never received. The workbook is the only witness.

This script reads the workbook and counts. It touches no database, writes
nothing, and needs no code change to run:

    python tools/check_lineage_fanin.py /path/to/legacy_lineage.xlsx

If it reports 0, the old lineage_id is losing nothing for this workbook and
step 4 of INSTALL_LINEAGE.md is optional in the fullest sense — the fix is
still correct, but nothing is currently being discarded.
"""
from __future__ import annotations
import os
import sys
import collections

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed — run this from the ingestion venv")

_PREFERRED = ("cplegacylineagesheet", "e2elineage", "lineage")


def _norm(s):
    return str(s or "").lower().replace(" ", "").replace("_", "")


def _pick(wb):
    norm = {_norm(s): s for s in wb.sheetnames}
    for p in _PREFERRED:
        if p in norm:
            return norm[p]
    for s in wb.sheetnames:
        if "lineage" in _norm(s) or "schema" in _norm(s):
            return s
    return wb.sheetnames[0]


def main(path, sheet=None):
    wb = load_workbook(path, data_only=True, read_only=True)
    name = sheet or _pick(wb)
    ws = wb[name]
    hdr = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        hdr = [str(c).strip() if c is not None else "" for c in row]
        break
    idx = {h: i for i, h in enumerate(hdr)}
    need = ("DWH_Target_Table", "DWH_Target_Column")
    missing = [c for c in need if c not in idx]
    if missing:
        sys.exit(f"sheet {name!r} has no {missing} column — pass the sheet name")

    def g(row, col):
        i = idx.get(col)
        if i is None or i >= len(row) or row[i] is None:
            return None
        v = str(row[i]).strip()
        return v or None

    rows = 0
    per_col = collections.Counter()
    src_of = collections.defaultdict(set)
    src_i = idx.get("SRC_Source_Table")
    for row in ws.iter_rows(min_row=2, values_only=True):
        t, c = g(row, "DWH_Target_Table"), g(row, "DWH_Target_Column")
        if not t or not c:
            continue
        rows += 1
        per_col[(t, c)] += 1
        if src_i is not None and src_i < len(row) and row[src_i] is not None:
            src_of[(t, c)].add(str(row[src_i]).strip())

    multi = {k: n for k, n in per_col.items() if n > 1}
    lost = sum(n - 1 for n in multi.values())

    print(f"sheet                     {name}")
    print(f"rows with a DWH target    {rows}")
    print(f"distinct target columns   {len(per_col)}")
    print(f"columns fed by 2+ rows    {len(multi)}")
    print(f"ROWS THE OLD KEY DISCARDS {lost}"
          f"   ({(100.0 * lost / rows):.1f}% of the sheet)" if rows else "")

    # A source file that only ever appears as a NON-surviving chain vanishes
    # from the load entirely — not undercounted, absent. The surviving row for
    # a key is simply the LAST one in sheet order, because that is what an
    # upsert on a duplicate id leaves behind.
    if src_i is not None:
        last_src, all_srcs = {}, set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            t, c = g(row, "DWH_Target_Table"), g(row, "DWH_Target_Column")
            if not t or not c:
                continue
            sv = g(row, "SRC_Source_Table")
            if not sv:
                continue
            all_srcs.add(sv)
            last_src[(t, c)] = sv          # later rows overwrite earlier ones
        survivors = set(last_src.values())
        gone = all_srcs - survivors
        print(f"\ndistinct source files in the sheet   {len(all_srcs)}")
        print(f"still present after the collapse     {len(survivors)}")
        print(f"SOURCE FILES THAT VANISH ENTIRELY    {len(gone)}")
        for sv in sorted(gone)[:15]:
            print(f"   {sv}")
        if len(gone) > 15:
            print(f"   … and {len(gone) - 15} more")

    if multi:
        print("\nworst-affected target columns:")
        for (t, c), n in sorted(multi.items(), key=lambda kv: -kv[1])[:10]:
            print(f"   {n:>2} chains   {t}.{c}")
    else:
        print("\nNo target column is fed by more than one row. The old "
              "lineage_id discards nothing for this workbook.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
