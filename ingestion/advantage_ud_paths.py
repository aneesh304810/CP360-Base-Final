"""One folder in, every AddVantage UD source found by what it IS.

The person loading the data drops every csv and workbook they have into
one folder (CP_ADDV_UD_DIR) and runs one step. File names vary between
drops, so each file is recognised by content first and by name second:

  the extract        a csv whose header has USER_DEFINED_ATTRIBUTE_CLOB
  profiler outputs   csvs with the profiler's exact names, in the folder
                     or in a profile/ sub-folder
  the UD workbook    an xlsx with sheets List and Tables, or a name that
                     contains "user defined"
  the TRP samples    an xlsx whose first sheet has "Account Number" and
                     "Ud <n> <seq>" headers, or a name that contains "trp"

A specific env var (CP_ADDV_UD_EXTRACT, CP_ADDV_UD_WORKBOOK, CP_ADDV_UD_TRP,
CP_ADDV_UD_PROFILE_DIR, CP_ADDV_UD_CODES) still wins over the scan, so a
file in an unusual place can be named explicitly.
"""
from __future__ import annotations

import csv
import os
import re
import sys

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))

PROFILE_FILES = ("attribute_profile.csv", "type_variance.csv", "parent_structures.csv", "schema_variants.csv",
                 "code_conflicts.csv", "run_summary.csv", "code_dictionary.csv")
CLOB = "USER_DEFINED_ATTRIBUTE_CLOB"


def default_dir():
    return os.environ.get("CP_ADDV_UD_DIR") or os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "local-data", "advantage-ud")


def _csv_header(path):
    try:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return [h.strip().upper() for h in next(csv.reader(fh))]
    except Exception:                                        # noqa: BLE001
        return []


def _xlsx_kind(path):
    """'workbook' | 'trp' | None, by sheet names and first-sheet headers."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        names = {ws.title.strip().lower() for ws in wb.worksheets}
        if "list" in names and "tables" in names:
            return "workbook"
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            vals = [str(c).strip().lower() for c in row if c is not None]
            if "account number" in vals and any(re.match(r"^ud\s+\d+", v) for v in vals):
                return "trp"
    except Exception:                                        # noqa: BLE001
        pass
    return None


def resolve(folder=None):
    """-> {"dir", "extract", "workbook", "trp", "profile_dir", "codes", "profile": {name: path},
           "unrecognised": [names]}"""
    d = folder or default_dir()
    out = {"dir": d, "extract": None, "workbook": None, "trp": None, "profile_dir": None,
           "codes": None, "profile": {}, "unrecognised": []}
    if not os.path.isdir(d):
        return out
    names = sorted(os.listdir(d))
    sub = os.path.join(d, "profile")
    for base in ([sub] if os.path.isdir(sub) else []) + [d]:
        for n in PROFILE_FILES:
            p = os.path.join(base, n)
            if os.path.exists(p) and n not in out["profile"]:
                out["profile"][n] = p
    if out["profile"]:
        out["profile_dir"] = os.path.dirname(next(iter(out["profile"].values())))
        out["codes"] = out["profile"].get("code_dictionary.csv")
    for n in names:
        p = os.path.join(d, n)
        low = n.lower()
        if not os.path.isfile(p) or low.startswith("~$") or n in ("README.md", ".gitignore"):
            continue
        if low in PROFILE_FILES:
            continue
        if low.endswith(".csv"):
            if CLOB in _csv_header(p) and out["extract"] is None:
                out["extract"] = p
            else:
                out["unrecognised"].append(n)
        elif low.endswith((".xlsx", ".xlsm")):
            kind = _xlsx_kind(p) or ("workbook" if "user defined" in low else "trp" if "trp" in low else None)
            if kind == "workbook" and out["workbook"] is None:
                out["workbook"] = p
            elif kind == "trp" and out["trp"] is None:
                out["trp"] = p
            else:
                out["unrecognised"].append(n)
        else:
            out["unrecognised"].append(n)
    # explicit variables win
    for key, var in (("extract", "CP_ADDV_UD_EXTRACT"), ("workbook", "CP_ADDV_UD_WORKBOOK"),
                     ("trp", "CP_ADDV_UD_TRP"), ("profile_dir", "CP_ADDV_UD_PROFILE_DIR"), ("codes", "CP_ADDV_UD_CODES")):
        if os.environ.get(var):
            out[key] = os.environ[var]
    return out


def describe(r):
    lines = [f"folder    {r['dir']}",
             f"extract   {r['extract'] or '- not found (a csv with ' + CLOB + ' in its header)'}",
             f"profile   {r['profile_dir'] or '- not found'}" + (f"  ({len(r['profile'])} of {len(PROFILE_FILES)} files)" if r["profile"] else ""),
             f"codes     {r['codes'] or '- not found (code_dictionary.csv)'}",
             f"workbook  {r['workbook'] or '- not found (sheets List + Tables)'}",
             f"trp       {r['trp'] or '- not found (Account Number + Ud n seq headers)'}"]
    if r["unrecognised"]:
        lines.append("ignored   " + ", ".join(r["unrecognised"]))
    return "\n".join(lines)
