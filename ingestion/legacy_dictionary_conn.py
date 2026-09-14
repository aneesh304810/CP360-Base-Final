"""
Legacy business dictionary connector — v2, built for the AddVantage master
workbook (LIST / ALL / DWH ALL / per-master sheets).

Reads the "ALL" sheet into legacy_dictionary and (optionally) the "DWH ALL"
sheet into legacy_lineage.

Workbook realities this handles (verified against the real file):
  - Row 1 is a title row; the header row is auto-detected (usually row 2).
  - Headers carry trailing spaces ("ADDV Field ", "Is Required ") -> stripped.
  - Field codes: PREFIX/NUMBER and PREFIX/NUMBER-SUB (BI/1, BI/2-1).
  - The DWH sheet mixes conventions for the SAME field: BI_2_L1 and BI/2-1.
    Join key = canonical code: separators -> "_", then "_L<digits>" -> "_<digits>",
    so BI/2-1 and BI_2_L1 both become BI_2_1.
  - Two blank-field-code section rows -> skipped.
  - Descriptions are multi-line, up to ~2.8k chars -> long_desc CLOB; the first
    line is extracted as short_desc.
  - DWH ALL has duplicate (target_table, target_column) pairs (one target,
    many sources) -> lineage_id is built on the full 4-column grain, never
    just target:column.
  - 128 DWH rows have blank SRC -> loaded as unmapped (src columns NULL).

Env:
  CP_LEGACY_DICT_XLSX          the workbook (single file, both sheets)
  CP_LEGACY_DICT_SYSTEM        source system label (default ADDVANTAGE)
  CP_LEGACY_DICT_SHEET         dictionary sheet name (default "ALL")
  CP_LEGACY_LINEAGE_SHEET      lineage sheet name (default "DWH ALL")
  CP_LEGACY_LINEAGE_FROM_XLSX  "1" to also load legacy_lineage from DWH ALL

  Per-system workbooks are still supported for future CRD / STAR files:
  CP_LEGACY_DICT_ADDVANTAGE / CP_LEGACY_DICT_CRD / CP_LEGACY_DICT_STAR

  v3 — multiple target warehouses (data sources):
  CP_LEGACY_SOURCES            "PBDW=/data/pbdw.xlsx;IMDS=/data/imds.xlsx"
                               one workbook per warehouse; each may contain a
                               dictionary sheet ("ALL"), a rich 24-column
                               lineage sheet, and/or the 4-column "DWH ALL"
                               sheet. Rich rows win; 4-col rows that describe
                               a (target, source) pair the rich sheet lacks
                               are appended as supplements.
  CP_LEGACY_DATA_SOURCE        warehouse tag for single-workbook mode
                               (default PBDW)
  CP_LEGACY_LINEAGE_SHEET_RICH rich lineage sheet name (default: auto-detect
                               the first sheet whose header row contains
                               "Functional_Group")
  Every legacy_lineage row is tagged with data_source and its lineage_id is
  prefixed with it: {ds}:{tgt}:{col}:{srchash}. Requires sql/29.
"""
from __future__ import annotations

import os
import re
import hashlib
import logging
from openpyxl import load_workbook

log = logging.getLogger("cp.legacy_dictionary")

SYSTEM_ENV = {
    "ADDVANTAGE": "CP_LEGACY_DICT_ADDVANTAGE",
    "CRD": "CP_LEGACY_DICT_CRD",
    "STAR": "CP_LEGACY_DICT_STAR",
}

# tolerant header matching: canonical field -> accepted header spellings
# (headers are stripped + lowercased + separators/parens collapsed first)
HEADER_ALIASES = {
    "field_code":        ["addv field", "field", "field code", "addvantage field code",
                          "crd field code", "star field code", "legacy field code", "code"],
    "field_name":        ["addv name", "name", "business term", "term", "asset name"],
    "master":            ["master"],
    "business_function": ["group", "business function", "function"],
    "data_type":         ["field data type", "data type", "type"],
    "max_length":        ["field max length", "max length", "length"],
    "num_precision":     ["precision if numeric", "precision"],
    "date_format":       ["format if date", "format"],
    "is_required":       ["is required", "required"],
    "is_unique":         ["is unique", "unique"],
    "long_desc":         ["description", "long description", "long desc"],
    "pb_field_mapping":  ["pb field mapping", "pb mapping"],
    "comments":          ["data selection comments", "data selection / comments",
                          "comments", "data selection"],
    "source_system":     ["source system", "system", "legacy system"],
    # legacy v1 sheets (classification-style) still resolve:
    "privacy_class":     ["data privacy classification", "privacy classification", "privacy"],
    "regulatory_class":  ["regulatory classification", "regulatory"],
    "operational_class": ["operational classification", "operational"],
    "status":            ["status"],
    "short_desc":        ["short description", "short desc"],
}

DWH_HEADERS = {
    "dwh_target_table":  ["dwh target table", "dwh_target_table"],
    "dwh_target_column": ["dwh target column", "dwh_target_column"],
    "src_source_table":  ["src source table", "src_source_table"],
    "src_source_column": ["src source column", "src_source_column"],
}

# rich (24-column) lineage sheet — full SRC->STG1->STG2->DWH chain.
# headers normalized by _norm_header first (separators/parens -> spaces).
RICH_HEADERS = {
    "dwh_target_table":        ["dwh target table"],
    "functional_group":        ["functional group"],
    "table_type":              ["table type"],
    "dwh_target_column":       ["dwh target column"],
    "dwh_type":                ["dwh type"],
    "dwh_length":              ["dwh length"],
    "dwh_precision":           ["dwh precision"],
    "stg2_source_table":       ["stg2 source table"],
    "stg2_source_column":      ["stg2 source column"],
    "stg2_to_dwh_transform":   ["stg2 to dwh transformation", "stg2 to dwh transform"],
    "stg2_type":               ["stg2 type"],
    "stg2_length":             ["stg2 length"],
    "stg2_precision":          ["stg2 precision"],
    "stg1_source_table":       ["stg1 source table"],
    "stg1_source_column":      ["stg1 source column"],
    "stg1_type":               ["stg1 type"],
    "stg1_length":             ["stg1 length"],
    "stg1_precision":          ["stg1 precision"],
    "src_source_table":        ["src source table"],
    "src_source_column":       ["src source column"],
    "src_to_stg1_transform":   ["src to stg1 transformation", "src to stg1 transform"],
    "stg1_to_stg2_transform":  ["stg1 to stg2 transformation", "stg1 to stg2 transform"],
    "lineage_status":          ["lineage status"],
    "lineage_status_detail":   ["lineage status detail"],
    "is_ud":                   ["is ud", "ud"],
    "ud_key":                  ["ud key"],
}


def _s(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _norm_header(h):
    # strip, lowercase, collapse separators AND parentheses:
    # "Precision(If Numeric)" -> "precision if numeric"
    return re.sub(r"[\s_\-./()]+", " ", str(h or "").strip().lower()).strip()


def normalize_code(code):
    """Separator normalization: BI/2-1 -> BI_2_1 ; ST.SEC.01 -> ST_SEC_01."""
    if not code:
        return None
    c = re.sub(r"[\s/.\-]+", "_", str(code).strip())
    c = re.sub(r"_{2,}", "_", c).strip("_")
    return c.upper()


def canonical_code(code):
    """The JOIN key. Normalizes separators, then collapses the DWH 'L' line
    convention: BI_2_L1 -> BI_2_1, so it meets the dictionary's BI/2-1.
    Safe because dictionary sub-numbers are always pure digits."""
    c = normalize_code(code)
    if not c:
        return None
    return re.sub(r"_L(\d+)", r"_\1", c)


def _first_line(text, limit=1990):
    if not text:
        return None
    line = str(text).strip().splitlines()[0].strip()
    return line[:limit] or None


def _find_header_row(ws, targets, scan=6):
    """Return (row_number, header_list) for the first row within `scan` rows
    whose cells contain one of the target header names (normalized)."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), start=1):
        normed = [_norm_header(c) for c in row]
        if any(h in targets for h in normed):
            return i, normed
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    return 1, [_norm_header(c) for c in first]


def _resolve(hdr, aliases):
    idx = {}
    for canon, names in aliases.items():
        for i, h in enumerate(hdr):
            if h in names:
                idx[canon] = i
                break
    return idx


class LegacyDictionaryConnector:
    name = "legacy_dictionary"

    def __init__(self, sources, dict_sheet=None, lineage_sheet=None, load_lineage=False,
                 rich_sheet=None):
        # sources: list of (source_system_or_None, xlsx_path, data_source)
        self.sources = [t if len(t) == 3 else (t[0], t[1], "PBDW") for t in sources]
        self.dict_sheet = dict_sheet or "ALL"
        self.lineage_sheet = lineage_sheet or "DWH ALL"
        self.rich_sheet = rich_sheet            # None -> auto-detect
        self.load_lineage = load_lineage

    @classmethod
    def from_env(cls):
        sources = []
        system = (os.environ.get("CP_LEGACY_DICT_SYSTEM") or "ADDVANTAGE").upper()
        default_ds = (os.environ.get("CP_LEGACY_DATA_SOURCE") or "PBDW").upper()

        # v3: one workbook per target warehouse — "PBDW=/a.xlsx;IMDS=/b.xlsx"
        multi = os.environ.get("CP_LEGACY_SOURCES")
        if multi:
            for part in multi.split(";"):
                part = part.strip()
                if not part or "=" not in part:
                    continue
                ds, path = part.split("=", 1)
                sources.append((system, path.strip(), ds.strip().upper() or "PBDW"))

        combined = os.environ.get("CP_LEGACY_DICT_XLSX")
        if combined:
            sources.append((system, combined, default_ds))
        for sysname, env in SYSTEM_ENV.items():
            path = os.environ.get(env)
            if path:
                sources.append((sysname, path, default_ds))
        if not sources:
            sources.append(("ADDVANTAGE",
                            "sample-artifacts/LEGACY-DICTIONARY/addvantage_master_workbook.xlsx",
                            default_ds))
        return cls(
            sources,
            dict_sheet=os.environ.get("CP_LEGACY_DICT_SHEET"),
            lineage_sheet=os.environ.get("CP_LEGACY_LINEAGE_SHEET"),
            load_lineage=os.environ.get("CP_LEGACY_LINEAGE_FROM_XLSX", "0") == "1",
            rich_sheet=os.environ.get("CP_LEGACY_LINEAGE_SHEET_RICH"),
        )

    # ------------------------------------------------------------------ parse
    def parse(self):
        dict_rows, lineage_rows = [], []
        for system, path, ds in self.sources:
            d, l = self._parse_one(system, path, ds)
            dict_rows.extend(d)
            lineage_rows.extend(l)
        log.info("legacy_dictionary: %d definitions, %d lineage rows parsed from %d source(s)",
                 len(dict_rows), len(lineage_rows), len(self.sources))
        return {"dict": dict_rows, "lineage": lineage_rows}

    def _parse_one(self, system, path, ds="PBDW"):
        if not os.path.exists(path):
            log.warning("legacy dictionary workbook not found: %s (skipping)", path)
            return [], []
        wb = load_workbook(path, data_only=True, read_only=True)

        # ---- dictionary sheet (default "ALL") — warehouse-agnostic; parsed
        # from any workbook that carries it, deduped downstream by dict_key ----
        dict_rows = []
        if self.dict_sheet in wb.sheetnames:
            dict_rows = self._parse_dict_sheet(wb[self.dict_sheet], system, path)
        elif not self.load_lineage:
            dict_rows = self._parse_dict_sheet(wb[wb.sheetnames[0]], system, path)

        # ---- lineage: rich 24-col sheet is primary, 4-col supplements ----
        lineage_rows = []
        if self.load_lineage:
            rich = []
            rich_ws = self._find_rich_sheet(wb)
            if rich_ws is not None:
                rich = self._parse_rich_lineage_sheet(rich_ws, path, ds)
            simple = []
            if self.lineage_sheet in wb.sheetnames:
                simple = self._parse_lineage_sheet(wb[self.lineage_sheet], path, ds)
            if not rich and not simple:
                log.warning("no lineage sheet (rich or '%s') found in %s",
                            self.lineage_sheet, path)
            lineage_rows = self._merge_lineage(rich, simple)
            log.info("legacy_dictionary: %s [%s] -> %d lineage rows "
                     "(%d rich, %d from 4-col sheet)",
                     system, ds, len(lineage_rows), len(rich), len(simple))
        return dict_rows, lineage_rows

    def _find_rich_sheet(self, wb):
        """The rich sheet is identified by its header containing
        Functional_Group — name it via CP_LEGACY_LINEAGE_SHEET_RICH or let
        this auto-detect (skipping the 4-col sheet)."""
        if self.rich_sheet:
            return wb[self.rich_sheet] if self.rich_sheet in wb.sheetnames else None
        for name in wb.sheetnames:
            if name == self.lineage_sheet:
                continue
            ws = wb[name]
            _, hdr = _find_header_row(ws, targets={"dwh target table"})
            if "functional group" in hdr:
                return ws
        return None

    @staticmethod
    def _lin_key(tgt_t, tgt_c, src_t, src_c):
        return (
            (tgt_t or "").strip().upper(),
            (tgt_c or "").strip().upper(),
            (normalize_code(src_t) or ""),
            (canonical_code(src_c) or ""),
        )

    def _merge_lineage(self, rich, simple):
        """Rich rows win; 4-col rows for (target, source) pairs the rich sheet
        does not know are appended as supplements."""
        out = {self._lin_key(r["dwh_target_table"], r["dwh_target_column"],
                             r.get("src_source_table"), r.get("src_source_column")): r
               for r in rich}
        added = 0
        for r in simple:
            k = self._lin_key(r["dwh_target_table"], r["dwh_target_column"],
                              r.get("src_source_table"), r.get("src_source_column"))
            if k not in out:
                r["lineage_status_detail"] = "From 4-column DWH ALL sheet (no staging detail)"
                out[k] = r
                added += 1
        if simple and rich:
            log.info("legacy_dictionary: 4-col sheet supplemented %d rows not in rich sheet", added)
        return list(out.values())

    def _parse_rich_lineage_sheet(self, ws, path, ds="PBDW"):
        hdr_row, hdr = _find_header_row(ws, targets={"dwh target table"})
        idx = _resolve(hdr, RICH_HEADERS)
        if "dwh_target_table" not in idx or "dwh_target_column" not in idx:
            log.warning("rich lineage sheet '%s' in %s missing target columns (skipping)",
                        ws.title, path)
            return []

        def g(row, canon):
            i = idx.get(canon)
            return _s(row[i]) if i is not None and i < len(row) else None

        out, seen = [], set()
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            tgt_t, tgt_c = g(row, "dwh_target_table"), g(row, "dwh_target_column")
            if not tgt_t or not tgt_c:
                continue
            src_t, src_c = g(row, "src_source_table"), g(row, "src_source_column")
            src_sig = hashlib.md5(f"{src_t or ''}|{src_c or ''}".encode()).hexdigest()[:12]
            lid = f"{ds}:{tgt_t}:{tgt_c}:{src_sig}"[:600]
            if lid in seen:
                continue
            seen.add(lid)
            is_ud = (g(row, "is_ud") or "").upper()
            out.append({
                "lineage_id": lid,
                "data_source": ds[:40],
                "dwh_target_table": tgt_t[:200],
                "dwh_target_column": tgt_c[:200],
                "dwh_type": g(row, "dwh_type"),
                "dwh_length": g(row, "dwh_length"),
                "dwh_precision": g(row, "dwh_precision"),
                "stg2_source_table": g(row, "stg2_source_table"),
                "stg2_source_column": g(row, "stg2_source_column"),
                "stg2_to_dwh_transform": g(row, "stg2_to_dwh_transform"),
                "stg2_type": g(row, "stg2_type"),
                "stg2_length": g(row, "stg2_length"),
                "stg2_precision": g(row, "stg2_precision"),
                "stg1_source_table": g(row, "stg1_source_table"),
                "stg1_source_column": g(row, "stg1_source_column"),
                "stg1_type": g(row, "stg1_type"),
                "stg1_length": g(row, "stg1_length"),
                "stg1_precision": g(row, "stg1_precision"),
                "src_source_table": src_t[:200] if src_t else None,
                "src_source_column": src_c[:200] if src_c else None,
                "src_to_stg1_transform": g(row, "src_to_stg1_transform"),
                "stg1_to_stg2_transform": g(row, "stg1_to_stg2_transform"),
                "lineage_status": g(row, "lineage_status") or ("Exists" if src_c else "unmapped"),
                "lineage_status_detail": g(row, "lineage_status_detail"),
                "is_ud": "Y" if is_ud.startswith("Y") else "N",
                "ud_key": g(row, "ud_key"),
                "functional_group": g(row, "functional_group"),
                "table_type": g(row, "table_type"),
            })
        log.info("legacy_dictionary: rich lineage sheet '%s' -> %d rows (header row %d)",
                 ws.title, len(out), hdr_row)
        return out

    def _parse_dict_sheet(self, ws, system, path):
        hdr_row, hdr = _find_header_row(ws, targets={"addv field", "field code", "code"})
        idx = _resolve(hdr, HEADER_ALIASES)
        if "field_code" not in idx:
            log.warning("legacy dictionary %s: no field-code column found in %s (skipping)",
                        system or "combined", path)
            return []

        def g(row, canon):
            i = idx.get(canon)
            return _s(row[i]) if i is not None and i < len(row) else None

        out = []
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            code = g(row, "field_code")
            if not code:            # blank-code section rows -> skip
                continue
            sys_val = (system or g(row, "source_system") or "ADDVANTAGE").strip().upper()
            canon = canonical_code(code)
            master = g(row, "master")
            desc = g(row, "long_desc")
            privacy = g(row, "privacy_class")
            master_key = (normalize_code(master) or "NA")[:40]
            out.append({
                "dict_key": f"{sys_val}:{canon}:{master_key}"[:200],
                "source_system": sys_val[:40],
                "field_code": code[:120],
                "field_code_norm": canon[:120],       # canonical join key
                "asset_name": g(row, "field_name"),
                "business_term": g(row, "field_name"),
                "business_function": g(row, "business_function"),
                "master_name": master,
                "data_type": g(row, "data_type"),
                "max_length": g(row, "max_length"),
                "num_precision": g(row, "num_precision"),
                "date_format": g(row, "date_format"),
                "is_required": ("Y" if (g(row, "is_required") or "").lower().startswith("y") else "N"),
                "is_unique": ("Y" if (g(row, "is_unique") or "").lower().startswith("y") else "N"),
                "short_desc": g(row, "short_desc") or _first_line(desc),
                "long_desc": desc,
                "pb_field_mapping": g(row, "pb_field_mapping"),
                "comments_txt": g(row, "comments"),
                "privacy_class": privacy,
                "regulatory_class": g(row, "regulatory_class"),
                "operational_class": g(row, "operational_class"),
                "status": g(row, "status"),
                "is_pii": "Y" if (privacy or "").strip().lower() in ("pii", "restricted", "sensitive") else "N",
            })
        log.info("legacy_dictionary: %s -> %d definitions from sheet '%s' (header row %d)",
                 system, len(out), ws.title, hdr_row)
        return out

    def _parse_lineage_sheet(self, ws, path, ds="PBDW"):
        hdr_row, hdr = _find_header_row(ws, targets={"dwh target table", "dwh_target_table"})
        idx = _resolve(hdr, DWH_HEADERS)
        if "dwh_target_table" not in idx or "dwh_target_column" not in idx:
            log.warning("DWH lineage sheet in %s missing target columns (skipping)", path)
            return []

        def g(row, canon):
            i = idx.get(canon)
            return _s(row[i]) if i is not None and i < len(row) else None

        out, seen = [], set()
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            tgt_t, tgt_c = g(row, "dwh_target_table"), g(row, "dwh_target_column")
            if not tgt_t or not tgt_c:
                continue
            src_t, src_c = g(row, "src_source_table"), g(row, "src_source_column")
            # composite grain: one target column can source from many files, so
            # the id must include the source side (380 duplicate target pairs).
            src_sig = hashlib.md5(f"{src_t or ''}|{src_c or ''}".encode()).hexdigest()[:12]
            lid = f"{ds}:{tgt_t}:{tgt_c}:{src_sig}"[:600]
            if lid in seen:          # exact duplicate row -> skip
                continue
            seen.add(lid)
            out.append({
                "lineage_id": lid,
                "data_source": ds[:40],
                "dwh_target_table": tgt_t[:200],
                "dwh_target_column": tgt_c[:200],
                "src_source_table": src_t[:200] if src_t else None,
                "src_source_column": src_c[:200] if src_c else None,
                "lineage_status": "mapped" if src_c else "unmapped",
                "is_ud": "N",
            })
        log.info("legacy_dictionary: DWH lineage sheet '%s' -> %d rows (header row %d)",
                 ws.title, len(out), hdr_row)
        return out

    # ------------------------------------------------------------------- load
    def load(self, loader, bundle):
        for r in bundle["dict"]:
            loader._merge("legacy_dictionary", ("dict_key",), r)
        for r in bundle.get("lineage", []):
            loader._merge("legacy_lineage", ("lineage_id",), r)
        loader.commit()
        log.info("legacy_dictionary: %d definitions + %d lineage rows loaded",
                 len(bundle["dict"]), len(bundle.get("lineage", [])))
