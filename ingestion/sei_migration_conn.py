"""SEI migration catalog connector: the merged source-file catalog, every
sheet of every workbook in one folder.

WHAT IT READS. Drop "SEI_All_Source_Files_Merged_Catalog_Expanded.xlsx"
(and any other workbook or csv shaped like it) into CP_SEI_MIGRATION_DIR.
Every sheet whose header row carries Catalog_ID and Source_Attribute is
read; other sheets are skipped and named in the log. Headers are matched
through _hkey(), which folds case, spaces, underscores, hyphens and
punctuation, so "BOXI Report & FieldNames" and "boxi_report_fieldnames"
are the same column. The two long DSR headers are matched by prefix.

WHAT IT WRITES (sql/78).
    cp_sei_migration_field    one row per catalog row, the sheet's columns
                              plus the readings from sei_migration_rules
    cp_sei_migration_source   the BBH tables and fields that feed the row
    cp_sei_migration_target   where the field goes next: outbound files,
                              BOXI reports, ADE-CAS, the desktop

Child rows for a catalog id are deleted and rewritten on every load, so a
table dropped from a rule does not linger in the lineage. Field rows are
MERGEd by catalog id. Set CP_SEI_MIGRATION_RELOAD=1 to empty all three
tables first (a row dropped from the workbook otherwise stays).

Run it:
    python -m ingestion.run sei_migration

Env:
    CP_SEI_MIGRATION_DIR     the folder (default local-data/sei-migration)
    CP_SEI_MIGRATION_RELOAD  "1" to replace rather than merge
"""
from __future__ import annotations

import csv
import logging
import os
import re
import sys

from .base import BaseConnector
from .sei_migration_rules import (classify_mandatory, classify_rule, country_specific, effective_rule,
                                  flags, normalize_status, parse_sources, parse_targets)

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
log = logging.getLogger("cp.sei_migration")

# header key -> field column. Keys are _hkey() of the sheet's headers; a key
# that is a prefix of the folded header also matches (the DSR columns).
HEADERS = {
    "CATALOGID": "catalog_id", "FUNCTIONALGROUP": "functional_group", "SOURCESYSTEM": "source_system",
    "SOURCEOBJECT": "source_object", "SEQUENCE": "seq", "SOURCEATTRIBUTE": "source_attribute",
    "DATATYPE": "data_type", "MAXLENGTH": "max_length", "MAXDECIMALPOS": "max_decimal",
    "FIRMDOMICILECOUNTRY": "domicile", "MANDATORY": "mandatory_text", "ADDITIONALVALIDATIONS": "validations",
    "REMARKS": "remarks", "ACCEPTABLEVALUES": "acceptable_values", "DESKTOP": "desktop",
    "BOXIREPORTFIELDNAMES": "boxi_text", "PROCESSINGLOGIC": "processing_logic",
    "STANDARDOUTBOUNDFILESFIELDNAMES": "outbound_field", "STANDARDOUTBOUNDFILESNAMES": "outbound_file",
    "STANDARDOUTBOUNDTRANSFORMATIONS": "outbound_transform", "ADECASCOMPONENTFIELDNAMES": "ade_cas",
    "FUNCTIONCATEGORY": "function_category",
    "COMMONLOGICEMPLOYEEDSR": "emp_dsr_logic", "EMPLOYEEDSRSTATUS": "emp_dsr_status",
    "COMMONLOGICTEAMDSR": "team_dsr_logic", "TEAMDSRSTATUS": "team_dsr_status",
    "OTHERMAPPINGLOGIC": "other_mapping_logic", "OTHERSTATUS": "other_status",
    "TABLESFIELDSOFFSYSTEM": "source_tables_text", "NOTE": "note", "NOTES": "note",
    "SOURCEWORKBOOK": "source_workbook", "SOURCESHEET": "source_sheet",
}
REQUIRED = ("CATALOGID", "SOURCEATTRIBUTE")
# column widths in sql/78; text is cut to fit rather than rejected
WIDTH = {"catalog_id": 120, "functional_group": 200, "source_system": 60, "source_object": 200, "source_attribute": 200,
         "data_type": 60, "max_length": 30, "max_decimal": 30, "domicile": 60, "mandatory_text": 4000, "validations": 4000,
         "remarks": 4000, "acceptable_values": 400, "desktop": 400, "boxi_text": 2000, "processing_logic": 4000,
         "outbound_field": 400, "outbound_file": 400, "outbound_transform": 2000, "ade_cas": 400, "function_category": 200,
         "emp_dsr_logic": 4000, "emp_dsr_status": 200, "team_dsr_logic": 4000, "team_dsr_status": 200,
         "other_mapping_logic": 4000, "other_status": 200, "source_tables_text": 2000, "note": 4000,
         "source_workbook": 400, "source_sheet": 200, "rule_text": 4000, "status_detail": 200, "lookup_name": 400}


def _hkey(h):
    return re.sub(r"[^A-Z0-9]", "", str(h or "").upper())


def _s(v):
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").strip()
    return "" if s.lower() in ("none", "nan") else s


def _cut(col, v):
    w = WIDTH.get(col)
    s = _s(v)
    return s[:w] if w and len(s) > w else s


def _int(v):
    try:
        return int(float(str(v).strip()))
    except Exception:                                        # noqa: BLE001
        return None


def map_headers(headers):
    """-> {column: index}. Exact folded match first, then prefix for the
    long DSR headers; a header no rule names is ignored."""
    idx = {}
    for i, h in enumerate(headers):
        k = _hkey(h)
        if not k:
            continue
        if k in HEADERS and HEADERS[k] not in idx:
            idx[HEADERS[k]] = i
            continue
        for key, col in HEADERS.items():
            if len(key) >= 12 and k.startswith(key) and col not in idx:
                idx[col] = i
                break
    return idx


def has_catalog_shape(headers):
    keys = {_hkey(h) for h in headers}
    return all(any(k.startswith(r) for k in keys) for r in REQUIRED)


def find_header_row(rows, scan=10):
    """(index, headers) of the first row among the first `scan` that has the
    catalog's required headers; else the row with the most text cells."""
    best, best_n = 0, -1
    for i, row in enumerate(rows[:scan]):
        vals = [_s(c) for c in row]
        if has_catalog_shape(vals):
            return i, vals
        n = sum(1 for v in vals if v and not re.fullmatch(r"[+-]?\d+(\.\d+)?", v))
        if n > best_n:
            best, best_n = i, n
    return best, [_s(c) for c in rows[best]] if rows else []


def read_row(idx, row, workbook, sheet):
    """One catalog row -> {"field": {...}, "sources": [...], "targets": [...]}."""
    get = lambda col: _s(row[idx[col]]) if col in idx and idx[col] < len(row) else ""     # noqa: E731
    f = {col: _cut(col, get(col)) for col in HEADERS.values() if col != "seq"}
    f["seq"] = _int(get("seq"))
    if not f.get("source_attribute"):
        return None
    if not f.get("catalog_id"):
        obj = re.sub(r"[^A-Z0-9]+", "-", (f.get("source_object") or "ROW").upper()).strip("-")
        f["catalog_id"] = f"SEI-{obj}-{(f['seq'] or 0):03d}-{re.sub(r'[^A-Z0-9]', '', f['source_attribute'].upper())[:40]}"[:120]
    if not f.get("source_workbook"):
        f["source_workbook"] = _cut("source_workbook", workbook)
    if not f.get("source_sheet"):
        f["source_sheet"] = _cut("source_sheet", sheet)
    # the readings
    m = classify_mandatory(f["mandatory_text"])
    f["mandatory_class"], f["mand_yes"], f["mand_no"], f["mand_na"] = m["class"], m["yes"], m["no"], m["na"]
    rule, rule_side = effective_rule(f["processing_logic"], f["other_mapping_logic"])
    f["rule_text"], f["rule_side"] = _cut("rule_text", rule), rule_side
    f["rule_class"] = classify_rule(rule, f["source_attribute"])
    f["status_class"], f["status_detail"] = normalize_status(f["other_status"])
    f["status_detail"] = _cut("status_detail", f["status_detail"])
    fl = flags(f["note"], f["other_mapping_logic"], f["remarks"], f["validations"])
    f.update(fl)
    f["country_specific"] = country_specific(f["domicile"])
    f["has_validation"] = "Y" if f["validations"] else "N"
    av = f["acceptable_values"]
    f["lookup_name"] = _cut("lookup_name", av if av else "")
    sources = parse_sources(f["source_tables_text"], rule, av)
    f["upstream_n"] = sum(1 for s in sources if s["role"] == "SOURCE")
    f["crosswalk_n"] = sum(1 for s in sources if s["role"] in ("CROSSWALK", "CONFIG"))
    f["systems"] = ",".join(sorted({s["system"] for s in sources if s["role"] == "SOURCE"}))[:400]
    targets = parse_targets(f["outbound_file"], f["outbound_field"], f["outbound_transform"], f["boxi_text"], f["ade_cas"], f["desktop"])
    cid = f["catalog_id"]
    src_rows = [{"catalog_id": cid, "source_table": s["table"][:200], "source_field": (s["field"] or "-")[:200],
                 "system_class": s["system"], "role": s["role"], "how": s["how"]} for s in sources]
    tgt_rows = [{"catalog_id": cid, "target_kind": t["kind"], "target_object": (t["object"] or "-")[:400],
                 "target_field": (t["field"] or "-")[:400], "transformation": (t["transformation"] or "")[:2000]} for t in targets]
    return {"field": f, "sources": src_rows, "targets": tgt_rows}


class SeiMigrationConnector(BaseConnector):
    name = "sei_migration"

    def __init__(self, folder=None, reload=False):
        self.folder = folder or os.environ.get("CP_SEI_MIGRATION_DIR") or os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "local-data", "sei-migration")
        self.reload = reload

    @classmethod
    def from_env(cls):
        return cls(reload=os.environ.get("CP_SEI_MIGRATION_RELOAD", "").strip() in ("1", "true", "yes"))

    # ------------------------------------------------------------- parse
    def files(self):
        if not os.path.isdir(self.folder):
            return []
        return sorted(os.path.join(self.folder, n) for n in os.listdir(self.folder)
                      if not n.startswith("~$") and n.lower().rsplit(".", 1)[-1] in ("xlsx", "xlsm", "csv"))

    def parse(self):
        b = {"fields": {}, "sources": [], "targets": [], "sheets": [], "skipped": [], "files": []}
        files = self.files()
        if not files:
            log.warning("sei_migration: no workbook or csv in %s (set CP_SEI_MIGRATION_DIR)", self.folder)
            return b
        for path in files:
            name = os.path.basename(path)
            b["files"].append(name)
            if name.lower().endswith(".csv"):
                with open(path, newline="", encoding="utf-8-sig") as fh:
                    rows = list(csv.reader(fh))
                self._take(b, rows, name, "csv")
                continue
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                self._take(b, rows, name, ws.title)
        b["sources"] = [s for cid in b["fields"] for s in b["fields"][cid]["sources"]]
        b["targets"] = [t for cid in b["fields"] for t in b["fields"][cid]["targets"]]
        log.info("sei_migration: %s fields from %s sheet(s); skipped %s", len(b["fields"]), len(b["sheets"]), b["skipped"] or "none")
        return b

    def _take(self, b, rows, workbook, sheet):
        rows = [r for r in rows if r and any(_s(c) for c in r)]
        if not rows:
            return
        hi, headers = find_header_row(rows)
        if not has_catalog_shape(headers):
            b["skipped"].append(f"{workbook}:{sheet}")
            return
        idx = map_headers(headers)
        n = 0
        for row in rows[hi + 1:]:
            r = read_row(idx, list(row), workbook, sheet)
            if r:
                b["fields"][r["field"]["catalog_id"]] = r
                n += 1
        b["sheets"].append({"workbook": workbook, "sheet": sheet, "rows": n, "columns": sorted(idx)})

    # -------------------------------------------------------------- load
    def load(self, loader, bundle=None):
        from .advantage_ud_bulk import bulk_merge
        b = bundle or self.parse()
        if not b["fields"]:
            return 0
        conn = getattr(loader, "conn", None)
        cur = conn.cursor() if conn is not None and hasattr(conn, "cursor") else None
        ids = list(b["fields"])
        if cur is not None:
            if self.reload:
                for t in ("cp_sei_migration_target", "cp_sei_migration_source", "cp_sei_migration_field"):
                    cur.execute(f"DELETE FROM {t}")
            else:
                for t in ("cp_sei_migration_target", "cp_sei_migration_source"):
                    cur.executemany(f"DELETE FROM {t} WHERE catalog_id = :1", [(i,) for i in ids])
            cur.close()
        fields = [b["fields"][i]["field"] for i in ids]
        n = bulk_merge(loader, "cp_sei_migration_field", ("catalog_id",), fields)
        bulk_merge(loader, "cp_sei_migration_source", ("catalog_id", "source_table", "source_field"), b["sources"])
        bulk_merge(loader, "cp_sei_migration_target", ("catalog_id", "target_kind", "target_object", "target_field"), b["targets"])
        loader.commit()
        log.info("sei_migration: %s fields, %s source links, %s target links", n, len(b["sources"]), len(b["targets"]))
        return n
