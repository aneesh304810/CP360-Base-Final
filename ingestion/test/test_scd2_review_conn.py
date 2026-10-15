"""The reviewed workbook -> scd2_rule_review, on a synthetic file.

    python ingestion/test/test_scd2_review_conn.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
from openpyxl import Workbook                                     # noqa: E402
from ingestion.scd2_review_conn import Scd2ReviewConnector        # noqa: E402
from tools.scd2_review_export import HEAD, write_workbook         # noqa: E402

BAD = 0


def ok(cond, msg, got=None):
    global BAD
    print(f"{'ok  ' if cond else 'FAIL'} {msg}" + ("" if cond else f"  -> {got!r}"))
    if not cond:
        BAD += 1


class FakeLoader:
    def __init__(self):
        self.rows, self.committed = [], False

    def _merge(self, table, pk, values, protect=()):
        self.rows.append((table, dict(values)))

    def commit(self):
        self.committed = True


with tempfile.TemporaryDirectory() as d:
    # the export, from rows shaped like the hub's query
    rows = [{"table_name": "DIM_ACCOUNT", "rule_id": "R2", "pattern": "MERGE_SCD2", "rule_text": "closes the prior version", "confidence": "PROVEN", "evidence": "step 14"},
            {"table_name": "DIM_ACCOUNT", "rule_id": "R8", "pattern": "MERGE_SCD2", "rule_text": "versions on real change", "confidence": "INFERRED", "evidence": "db"},
            {"table_name": "DIM_SECURITY", "rule_id": "R1", "pattern": "MERGE_SCD2", "rule_text": "start date", "confidence": "PROVEN", "evidence": "var",
             "business_meaning": "already written", "confirmed": "Y", "reviewer": "C. Prior"}]
    p = write_workbook(rows, os.path.join(d, "review.xlsx"))
    from openpyxl import load_workbook
    ws = load_workbook(p)["rules"]
    ok([c.value for c in ws[1]] == HEAD and ws.max_row == 4, "the export has the header and one row per rule", ws.max_row)
    ok(ws["G4"].value == "already written" and ws["H4"].value == "Y", "an earlier sign-off is pre-filled")
    ok(ws["D2"].protection.locked and not ws["G2"].protection.locked and ws.protection.sheet, "the tool's columns are locked, the reviewer's are open")

    # the reviewers fill it
    ws["G2"], ws["H2"], ws["I2"] = "Yesterday's record is closed the day before the new one starts.", "Y", "A. Reviewer"
    ws["H3"], ws["I3"], ws["J3"] = "N", "B. Reviewer", "one feed re-sends unchanged rows"
    load_workbook(p).save(p)  # no-op save keeps the file valid
    wb = load_workbook(p); wb["rules"]["G2"] = ws["G2"].value; wb["rules"]["H2"] = "Y"; wb["rules"]["I2"] = "A. Reviewer"
    wb["rules"]["H3"] = "N"; wb["rules"]["I3"] = "B. Reviewer"; wb["rules"]["J3"] = "one feed re-sends unchanged rows"
    wb.save(p)

    c = Scd2ReviewConnector(p)
    b = c.parse()
    ok(len(b["rows"]) == 3 and b["skipped"] == 0, "three reviewed rows, none skipped", (len(b["rows"]), b["skipped"]))
    L = FakeLoader(); n = c.load(L, b)
    ok(n == 3 and L.committed and all(t == "scd2_rule_review" for t, _ in L.rows), "all merged into the sign-off table")
    r2 = next(v for t, v in L.rows if v["rule_id"] == "R2")
    ok(r2["confirmed"] == "Y" and r2["reviewer"] == "A. Reviewer" and r2["rule_text_seen"] == "closes the prior version" and r2["source_file"] == "review.xlsx",
       "a confirmation carries the reviewer, the rule text seen and the file", r2)

    # the rules: no reviewer -> skipped; N without a note -> refused; bad CONFIRMED -> refused
    wb = load_workbook(p); wb["rules"]["I2"] = None; wb.save(p)
    ok(Scd2ReviewConnector(p).parse()["skipped"] == 1, "a row without a reviewer is not loaded")
    wb = load_workbook(p); wb["rules"]["I2"] = "A. Reviewer"; wb["rules"]["J3"] = None; wb.save(p)
    try:
        Scd2ReviewConnector(p).parse(); ok(False, "N without a note is refused")
    except ValueError as e:
        ok("needs a REVIEWER_NOTE" in str(e), "N without a note is refused", str(e))
    wb = load_workbook(p); wb["rules"]["J3"] = "why"; wb["rules"]["H2"] = "maybe"; wb.save(p)
    try:
        Scd2ReviewConnector(p).parse(); ok(False, "CONFIRMED other than Y/N is refused")
    except ValueError as e:
        ok("must be Y or N" in str(e), "CONFIRMED other than Y/N is refused", str(e))
    ok(Scd2ReviewConnector(os.path.join(d, "missing.xlsx")).parse()["rows"] == [], "a missing workbook is nothing to load, not a crash")

print(f"\n{BAD} assertion(s) failed" if BAD else "\nscd2-review assertions pass")
sys.exit(1 if BAD else 0)
