"""The three source files, ingested from synthetic copies shaped like them.

WHY THESE ASSERTIONS. The extract's CLOB has newlines inside quotes and
account numbers that must stay strings; the workbook's List sheet is a
multi-block layout and its Tables sheet hands codes back as floats; the
TRP samples carry names and must leave nothing behind but counts. Each
of those is a way a naive load silently corrupts the catalogue, and each
is pinned here. Then the brief's quality rules: invalid JSON is
quarantined and the load goes on, an unknown key is registered, a new key
set is recorded, and a missing attribute is never a deletion.

    python api/test/test_advantage_ud_ingest.py
"""
import csv
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from ingestion.advantage_ud_parser import parse_payload, classify_value, parse_key   # noqa: E402
from ingestion.advantage_ud_extract_conn import AdvantageUdExtractConnector         # noqa: E402
from ingestion.advantage_ud_workbook_conn import (AdvantageUdWorkbookConnector,      # noqa: E402
                                                  trp_header_to_key, find_header_row)
from ingestion.advantage_ud_trp_conn import AdvantageUdTrpConnector                 # noqa: E402
from ingestion.advantage_ud_bulk import merge_sql, bulk_merge                       # noqa: E402

bad = 0


def ok(cond, msg, got=None):
    global bad
    print(("ok   " if cond else "FAIL ") + msg + ("" if cond else f"  -> {got!r}"))
    if not cond:
        bad += 1


class Store:
    """Fake loader: tables of dicts keyed by pk, same semantics as MERGE."""
    def __init__(self):
        self.t = {}; self.commits = 0
    def _merge(self, table, pk, v):
        k = tuple(v[c] for c in pk)
        self.t.setdefault(table, {})[k] = {**self.t.get(table, {}).get(k, {}), **v}
    def commit(self):
        self.commits += 1
    def rows(self, table):
        return list(self.t.get(table, {}).values())


print("-- the parser agrees with the profiler's examples")
p = parse_payload('{"UD_1":"2=CLIENT ACCOUNT","UD_14":"TE=TAX EXEMPT","UD_23_1":"000007447","UD_23_2":"WHITMAN",'
                  '"UD_32_1":"MANUAL BILLING","UD_503":"S=SINGLY"}')
by = {a["attribute_name"]: a for a in p["attributes"]}
ok(by["UD_1"]["value_type"] == "CODE_DESCRIPTION" and by["UD_1"]["code_value"] == "2" and by["UD_1"]["description_value"] == "CLIENT ACCOUNT", "code=description split")
ok(by["UD_23_1"]["value_type"] == "IDENTIFIER_LEADING_ZERO" and by["UD_23_1"]["raw_value"] == "000007447" and by["UD_23_1"]["leading_zero_ind"] == "Y", "leading zeros kept and flagged")
ok(by["UD_23_1"]["parent_attribute"] == "UD_23" and by["UD_23_1"]["sequence_number"] == 1 and by["UD_23_1"]["key_structure"] == "MULTIPART", "multipart key split")
ok(by["UD_1"]["parent_attribute"] is None and by["UD_1"]["key_structure"] == "SINGLE", "a single has no parent in the catalogue")
ok(by["UD_32_1"]["value_type"] == "TEXT" and by["UD_503"]["code_value"] == "S", "text stays text, a one-letter code splits")
ok(len(p["schema_signature"]) == 16 and p["schema_signature"].isupper() and p["typed_signature"] != p["schema_signature"], "signatures are the profiler's 16 upper hex")
ok(classify_value("ANNUAL MIN = 87500")["code_value"] == "ANNUAL MIN", "the parser splits prose like the profiler; the classifier, not the parser, rejects it later")
ok(classify_value("07/03/2002")["date_mask"] == "MM/DD/YYYY|DD/MM/YYYY", "an ambiguous date keeps both masks")
ok(classify_value("1010000017")["value_type"] == "INTEGER" and classify_value("N/A")["value_type"] == "NULL_LITERAL" and classify_value("")["value_type"] == "BLANK", "integers, null literals, blanks")
ok(parse_payload('"{""UD_613"":""Y""}"')["attributes"][0]["value_type"] == "BOOLEAN_FLAG", "the doubly-quoted CSV form is unwrapped")
ok(parse_payload("{oops")["error"].startswith("INVALID_JSON") and parse_payload("[1]")["error"] == "TOP_LEVEL_LIST_NOT_OBJECT" and parse_payload("")["error"] == "EMPTY_PAYLOAD", "errors are named")
ok(parse_key("ACCOUNT")["key_structure"] == "NON_STANDARD" and parse_key("ACCOUNT")["is_ud_attribute"] == "N", "a non-UD key is NON_STANDARD")

print("-- the extract, streamed")
d = tempfile.mkdtemp()
csvp = os.path.join(d, "dataVar.csv")
with open(csvp, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh, quoting=csv.QUOTE_ALL)
    w.writerow(["ACCOUNT_NUMBER", "AS_OF_DATE", "USER_DEFINED_ATTRIBUTE_CLOB", "LOAD_DATE", "BATCH_ID", "LOAD_TYPE",
                "FIS_LOAD_DATE", "ACTIVE_IND", "ACCOUNT_UD_KEY", "ACCOUNT_KEY", "CREATED_TSP", "UPDATED_TSP"])
    base = ["07-OCT-26", None, "07-OCT-26 03.26.39.571109000 AM", "20261007032517992253", "EOD", "07-OCT-26 03.00.00.000000000 AM", "", None, None,
            "07-OCT-26 03.27.54.553848000 AM", "07-OCT-26 03.27.54.553848000 AM"]
    def row(acct, key, akey, clob):
        r = ["0" + acct if len(acct) < 10 else acct] + base[:]; r[2] = clob; r[8] = key; r[9] = akey; return r
    w.writerow(row("5091002413", "35841747", "462644", '{"UD_1":"2=CLIENT ACCOUNT",\n "UD_14":"TE=TAX EXEMPT",\n "UD_23_1":"000007447",\n "UD_23_2":"PLACEHOLDER",\n "UD_613":"Y"}'))
    w.writerow(row("5091002414", "35841748", "462645", '{"UD_1":"6=PARTNER ACCOUNT","UD_613":"Y","UD_999":"NEW"}'))
    w.writerow(row("5091002415", "35841749", "462646", '{"UD_613":"Y"}'))
    w.writerow(row("5091002416", "35841750", "462647", '{"UD_1":"2=CLIENT ACCOUNT",'))        # invalid JSON
    w.writerow(row("5091002417", "", "462648", '{"UD_613":"N"}'))                              # no key
c = AdvantageUdExtractConnector(csvp, chunk=2, known_attributes={"UD_1", "UD_14", "UD_23_1", "UD_23_2", "UD_613"})
L = Store()
out = c.load(L)
ok(out["rows"] == 5 and out["ok"] == 3 and out["quarantined"] == 2, "five rows: three loaded, two quarantined", out)
raws = {r["account_ud_key"]: r for r in L.rows("cp_advantage_ud_raw")}
ok(len(raws) == 4 and raws[35841750]["parse_status"] == "QUARANTINED", "the bad-JSON row keeps its raw row, marked", list(raws))
ok(raws[35841747]["batch_id"] == "20261007032517992253" and isinstance(raws[35841747]["batch_id"], str), "the 20-digit batch id is a string")
ok(raws[35841747]["key_count"] == 5 and raws[35841747]["payload_hash"] and raws[35841747]["schema_signature"], "hash, key count and signature on the raw row")
ok(all(not k.startswith("_") for r in L.rows("cp_advantage_ud_raw") for k in r), "no helper field reaches the table")
attrs = {a["attr_key"]: a for a in L.rows("cp_advantage_ud_attribute")}
ok(out["attributes"] == 9 and len(attrs) == 9, "one attribute row per key of a good row", out["attributes"])
ok(attrs["35841747:UD_23_1"]["raw_value"] == "000007447" and attrs["35841747:UD_23_1"]["value_type"] == "IDENTIFIER_LEADING_ZERO", "the embedded-newline CLOB parsed, zeros kept")
ok(attrs["35841747:UD_23_1"]["account_number"] == "5091002413", "the account number is the string it was")
q = {x["quarantine_key"]: x for x in L.rows("cp_advantage_ud_quarantine")}
ok("35841750" in q and q["35841750"]["parse_error"].startswith("INVALID_JSON") and q["35841750"]["payload_sample"].startswith('{"UD_1"'), "invalid JSON quarantined with a sample")
ok("row:6" in q and q["row:6"]["parse_error"] == "MISSING_ACCOUNT_UD_KEY", "a row without its key is quarantined by row number")
reg = {r["attribute_name"]: r for r in L.rows("cp_advantage_ud_registry")}
ok(list(reg) == ["UD_999"] and reg["UD_999"]["class_source"] == "INFERRED" and reg["UD_999"]["key_structure"] == "SINGLE", "the unknown key is registered, not refused", list(reg))
sch = L.rows("cp_advantage_ud_schema")
ok(len(sch) == 3 and out["new_schemas"] == 3 and all(s["attribute_list"] for s in sch), "three new exact key sets recorded with their key lists", [s["attribute_list"] for s in sch])
ok(L.commits >= 3, "flushed in chunks and committed")
# missing attribute is not a deletion: load again with UD_14 gone from the first row
with open(csvp, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh, quoting=csv.QUOTE_ALL)
    w.writerow(["ACCOUNT_NUMBER", "AS_OF_DATE", "USER_DEFINED_ATTRIBUTE_CLOB", "LOAD_DATE", "BATCH_ID", "LOAD_TYPE",
                "FIS_LOAD_DATE", "ACTIVE_IND", "ACCOUNT_UD_KEY", "ACCOUNT_KEY", "CREATED_TSP", "UPDATED_TSP"])
    w.writerow(row("5091002413", "35841747", "462644", '{"UD_1":"2=CLIENT ACCOUNT","UD_613":"Y"}'))
AdvantageUdExtractConnector(csvp, known_attributes={"UD_1", "UD_613"}).load(L)
ok("35841747:UD_14" in {a["attr_key"] for a in L.rows("cp_advantage_ud_attribute")}, "a key absent from a later payload is not deleted")
ok(AdvantageUdExtractConnector(os.path.join(d, "nope.csv")).load(L)["rows"] == 0, "a missing file loads nothing and does not raise")

print("-- bulk merge")
sql = merge_sql("cp_advantage_ud_attribute", ("attr_key",), ["attr_key", "raw_value"])
ok("MERGE INTO cp_advantage_ud_attribute t USING (SELECT :attr_key AS attr_key, :raw_value AS raw_value FROM dual) s ON (t.attr_key = s.attr_key)" in sql
   and "UPDATE SET t.raw_value = s.raw_value" in sql, "the bulk statement is the loader's MERGE", sql)
class _Cur:
    def __init__(self): self.calls = []
    def executemany(self, sql, rows): self.calls.append(len(rows))
    def close(self): pass
class _Conn:
    def __init__(self): self.c = _Cur()
    def cursor(self): return self.c
class _L:
    def __init__(self): self.conn = _Conn()
    def _merge(self, *a): raise AssertionError("should not be called with a connection")
l2 = _L(); n = bulk_merge(l2, "t", ("k",), [{"k": i} for i in range(12)], batch=5)
ok(n == 12 and l2.conn.c.calls == [5, 5, 2], "batched executemany on the loader's connection", l2.conn.c.calls)

print("-- the UD workbook")
from openpyxl import Workbook
wbp = os.path.join(d, "AddV User Defined Fields test.xlsx")
wb = Workbook(); ws = wb.active; ws.title = "List"
ws.append(["AddVantage User Defined Fields"])          # title row
ws.append(["UD Fields", None, "Tables", "Rubal Notes", "Type", "Name", "Type Description"])
ws.append(["1=ACCOUNTMASTER", 224, "6 - TABLE 714 - TAX AND TAX EXEMPT", "Tax status", 1, "Date", "Accepts MMDDYY or MM/DD/YY"])
ws.append(["3=PARTY", 13, "6 - TABLE 5 - OFFICER TABLE", "Client AML Risk and Politically exposed Client", 2, "Yes/No", "Y or N"])
ws.append(["34=MASTERACCOUNT", 0, "6 - TABLE 717 - OWNED BY", None, 3, "Text", "Allows definition of up to 99 lines of 32 characters each."])
ws.append([None, None, None, None, 4, "Money", "Up to 10 digits"])
ws.append([None, None, None, None, 6, "Table", "Value must come from a system table"])
ws2 = wb.create_sheet("Tables")
ws2.append(["Id", "Table Number", "Table Name", "Code", "Description1", "Description2", "Description3", "Description4", "Description5", "Rubal Notes"])
for i, (tn, name, code, desc) in enumerate([(714, "TAX AND TAX EXEMPT", "TE", "TAX EXEMPT"), (714, "TAX AND TAX EXEMPT", "TX", "TAXABLE"),
                                             (717, "OWNED BY", 1, "INTERNAL/FIRM ACCOUNT"), (717, "OWNED BY", 2, "CLIENT ACCOUNT"), (717, "OWNED BY", 6, "PARTNER ACCOUNT"),
                                             (5, "OFFICER TABLE", "000100", "A PERSON"), (5, "OFFICER TABLE", 1000, "ANOTHER PERSON")]):
    ws2.append([i + 1, tn, name, code, desc, None, None, None, None, None])
wb.save(wbp)
observed = {"UD_1": {"1", "2", "3", "4", "5", "6", "7"}, "UD_14": {"TE", "TX"}, "UD_503": {"S", "J"}}
cw = AdvantageUdWorkbookConnector(wbp, observed)
b = cw.parse()
ok(b["sheets"] == ["List", "Tables"] and b["mapping_sheet"] is None, "sheets found, no mapping sheet")
tabs = {t["table_number"]: t for t in b["tables"]}
ok(set(tabs) == {714, 717, 5} and tabs[714]["table_type"] == 6 and tabs[714]["table_name"] == "TAX AND TAX EXEMPT", "the List catalogue is split into type, number, name", tabs)
ok(tabs[5]["is_pii"] == "Y" and tabs[714]["is_pii"] == "N" and tabs[5]["code_count"] == 2, "table 5 is PII; code counts come from the Tables sheet")
ok(tabs[5]["sme_notes"].startswith("Client AML"), "SME notes kept as evidence")
ok([f["type_code"] for f in b["field_types"]] == [1, 2, 3, 4, 6] and b["field_types"][2]["value_class"] == "TEXT", "the five field types with their classes")
ok([(e["entity_code"], e["ud_count"]) for e in b["entities"]] == [(1, 224), (3, 13), (34, 0)], "entities with their counts from the unlabeled column", b["entities"])
codes = {(c["table_number"], c["code_value"]): c for c in b["codes"]}
ok(codes[(717, "1")]["code_value"] == "1" and codes[(5, "000100")]["code_value"] == "000100" and codes[(717, "6")]["description_value"] == "PARTNER ACCOUNT",
   "a numeric cell is a code string again, a leading-zero code survives", list(codes))
ok(codes[(5, "000100")]["description_value"] == "•••" and codes[(5, "000100")]["is_pii"] == "Y", "officer names never leave the workbook")
links = {l["attribute_name"]: l for l in b["links"]}
ok(links["UD_14"]["table_number"] == 714 and links["UD_14"]["link_status"] == "STRONGLY_INFERRED" and links["UD_14"]["overlap_pct"] == 100.0, "UD_14 -> table 714 inferred from full overlap", links.get("UD_14"))
ok(links["UD_1"]["table_number"] == 717 and links["UD_1"]["link_status"] == "WEAK" and links["UD_1"]["overlap_pct"] < 80, "3 of 7 observed codes is a WEAK link, not a strong one", links.get("UD_1"))
ok("UD_503" not in links, "no table shares a code: no link, no guess")
ok(b["notes"] == ["no UD->table sheet; links inferred from code overlap"], "the inference is said")
LW = Store(); nw = cw.load(LW, b)
dic = {k: v for k, v in LW.t["cp_advantage_ud_dictionary"].items()}
ok(("UD_14:TE:TABLES",) in dic and dic[("UD_14:TE:TABLES",)]["link_status"] == "STRONGLY_INFERRED", "a strongly linked table's codes are filed under the UD key")
ok(("UD_T717:2:TABLES",) in dic and ("UD_1:2:TABLES",) not in dic, "a weak link leaves the table's codes under the table, not the key")
ok(len(LW.rows("cp_advantage_ud_link")) == 2 and len(LW.rows("cp_advantage_ud_table")) == 3 and LW.commits == 1, "links and tables loaded, one commit")
# a mapping sheet, when present, wins
wb3 = Workbook(); wb3.remove(wb3.active)
wb3.create_sheet("List").append(["UD Fields", None, "Tables", "Rubal Notes", "Type", "Name", "Type Description"])
m = wb3.create_sheet("UD to Table"); m.append(["UD", "Table Number"]); m.append(["UD/1", 717]); m.append(["UD/14", "TABLE 714"])
wbp3 = os.path.join(d, "AddV User Defined Fields map.xlsx"); wb3.save(wbp3)
b3 = AdvantageUdWorkbookConnector(wbp3, observed).parse()
ok(b3["mapping_sheet"] == "UD to Table" and {l["attribute_name"]: l["link_status"] for l in b3["links"]} == {"UD_1": "VERIFIED", "UD_14": "VERIFIED"},
   "a sheet that maps UD to table makes the links VERIFIED", b3["links"])
ok(AdvantageUdWorkbookConnector(os.path.join(d, "none.xlsx")).parse()["tables"] == [], "a missing workbook loads nothing")

print("-- TRP reconciliation, counts only")
ok(trp_header_to_key("Ud 23 1") == "UD_23_1" and trp_header_to_key("Ud 514 10") == "UD_514_10" and trp_header_to_key("Ud 503") == "UD_503" and trp_header_to_key("Account Number") is None, "header rule")
trp = os.path.join(d, "TRP AddV UD MultiLine test.xlsx")
wt = Workbook(); wst = wt.active
wst.append(["Account Number", "Ud 23 1", "Ud 23 2", "Ud 514 1", "Ud 514 2", "Notes"])
wst.append([5091002413, "000007447", "PLACEHOLDER", "TRU", "ST"])
wst.append([5091002414, "N/A", "DELETE FILE", "", "SILENT"])
wst.append([9999999999, "000000001", "X", "", ""])
wt.save(trp)
ext = {"5091002413": {"UD_23_1": "000007447", "UD_23_2": "placeholder", "UD_1": "2=CLIENT ACCOUNT"},
       "5091002414": {"UD_23_2": "OTHER"}}
ct = AdvantageUdTrpConnector(trp, lambda a: ext.get(a), registry_keys={"UD_23_1", "UD_23_2"})
bt = ct.parse()
rc = {r["attribute_name"]: r for r in bt["recon"]}
ok(bt["accounts"] == 3 and bt["columns"] == 4 and bt["undocumented"] == ["Notes"], "accounts, UD columns, and the undocumented column named", (bt["accounts"], bt["columns"], bt["undocumented"]))
ok(rc["UD_23_1"]["value_equal"] == 1 and rc["UD_23_1"]["null_tokens"] == 1 and rc["UD_23_1"]["account_not_in_extract"] == 1, "a matching id, a null token, an unknown account", rc["UD_23_1"])
ok(rc["UD_23_2"]["value_equal"] == 1 and rc["UD_23_2"]["value_differs"] == 1 and rc["UD_23_2"]["junk_markers"] == 1, "case-insensitive equality, one difference, one junk marker", rc["UD_23_2"])
ok(rc["UD_514_1"]["missing_in_extract"] == 1 and rc["UD_514_1"]["in_registry"] == "N", "UD_514 is in the samples and not in the extract: the finding, as a number", rc["UD_514_1"])
ok(all(k in ("attribute_name", "trp_rows", "matched_accounts", "value_equal", "value_differs", "missing_in_extract",
             "account_not_in_extract", "null_tokens", "junk_markers", "in_registry") for r in bt["recon"] for k in r), "nothing but counts is stored")
LT = Store(); ok(ct.load(LT, bt) == 4 and LT.commits == 1, "four attributes reconciled, one commit")

print("-- one folder, every source recognised by content")
from ingestion.advantage_ud_paths import resolve, describe
import shutil
f = tempfile.mkdtemp()
shutil.copy(csvp, os.path.join(f, "DIM_ACCOUNT_UD extract oct.csv"))          # not called dataVar
shutil.copy(wbp, os.path.join(f, "metadata_v2.xlsx"))                        # no "user defined" in the name
shutil.copy(trp, os.path.join(f, "samples.xlsx"))                            # no "trp" in the name
with open(os.path.join(f, "attribute_profile.csv"), "w") as fh: fh.write("attribute_name\n")   # loose, not in profile/
with open(os.path.join(f, "code_dictionary.csv"), "w") as fh: fh.write("attribute_name,code_value,description_value\n")
with open(os.path.join(f, "notes.txt"), "w") as fh: fh.write("x")
with open(os.path.join(f, "other.csv"), "w") as fh: fh.write("a,b\n")
for v in ("CP_ADDV_UD_EXTRACT", "CP_ADDV_UD_WORKBOOK", "CP_ADDV_UD_TRP", "CP_ADDV_UD_PROFILE_DIR", "CP_ADDV_UD_CODES"):
    os.environ.pop(v, None)
r = resolve(f)
ok(r["extract"].endswith("DIM_ACCOUNT_UD extract oct.csv"), "the extract is the csv with the CLOB header, whatever its name", r["extract"])
ok(r["workbook"].endswith("metadata_v2.xlsx"), "the workbook is the xlsx with List and Tables sheets", r["workbook"])
ok(r["trp"].endswith("samples.xlsx"), "the samples are the xlsx with Account Number and Ud n seq headers", r["trp"])
ok(r["profile_dir"] == f and r["codes"].endswith("code_dictionary.csv") and set(r["profile"]) == {"attribute_profile.csv", "code_dictionary.csv"},
   "profiler csvs are found loose in the folder", (r["profile_dir"], list(r["profile"])))
ok(sorted(r["unrecognised"]) == ["notes.txt", "other.csv"], "anything else is listed as ignored, not guessed", r["unrecognised"])
os.environ["CP_ADDV_UD_EXTRACT"] = "/elsewhere/x.csv"
ok(resolve(f)["extract"] == "/elsewhere/x.csv", "a specific variable still wins over the scan")
os.environ.pop("CP_ADDV_UD_EXTRACT")
os.environ["CP_ADDV_UD_DIR"] = f
ok(resolve()["dir"] == f and AdvantageUdExtractConnector.from_env().path == r["extract"]
   and AdvantageUdWorkbookConnector.from_env().path == r["workbook"] and AdvantageUdTrpConnector.from_env().path == r["trp"],
   "CP_ADDV_UD_DIR drives every connector's default")
os.environ.pop("CP_ADDV_UD_DIR")
ok(resolve(os.path.join(f, "missing"))["extract"] is None, "a missing folder resolves to nothing, without raising")
ok("ignored   notes.txt, other.csv" in describe(r) and "extract   " in describe(r), "describe() prints what the step will log", describe(r))
sub = tempfile.mkdtemp(); os.makedirs(os.path.join(sub, "profile")); shutil.copy(os.path.join(f, "attribute_profile.csv"), os.path.join(sub, "profile", "attribute_profile.csv"))
ok(resolve(sub)["profile_dir"] == os.path.join(sub, "profile"), "a profile/ sub-folder is preferred when present")

print("-- run.py")
src = open(os.path.join(os.path.dirname(__file__), "..", "..", "ingestion", "run.py")).read()
ok(all(f'"{s}"' in src for s in ("advantage_ud_extract", "advantage_ud_workbook", "advantage_ud_trp", "advantage_ud")), "the steps are listed and dispatched")
i = src.index('if step == "advantage_ud":')
body = src[i:i + 3000]
ok(body.index("advantage_ud_profile") < body.index("advantage_ud_dictionary") < body.index("advantage_ud_extract") < body.index("advantage_ud_workbook") < body.index("advantage_ud_trp"),
   "the one step runs the five in dependency order")
ok("skipped, its source is not in the folder" in body and "inventory" in body, "it skips what is absent and writes the inventory first")
ps = open(os.path.join(os.path.dirname(__file__), "..", "..", "local", "load.ps1")).read()
ok('CP_ADDV_UD_DIR' in ps and '$Step -like "advantage_ud*"' in ps and 'takes a FOLDER' in ps, "load.ps1 knows the step, the variable and that -File is a folder")
ok("CP_ADDV_UD_DIR" in open(os.path.join(os.path.dirname(__file__), "..", "..", "local", "set-env.ps1")).read(), "set-env.ps1 sets the variable")

print()
print("advantage_ud_ingest assertions " + ("FAIL" if bad else "pass"))
if bad:
    sys.exit(1)
